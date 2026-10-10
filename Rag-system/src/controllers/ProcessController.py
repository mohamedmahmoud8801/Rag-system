from .BaseController import BaseController
from .ProjectController import ProjectController

import os
import io
import base64
import logging
import re

from pathlib import Path
from typing import List
from dataclasses import dataclass

from PIL import Image, ImageOps
from openpyxl import load_workbook

from langchain_community.document_loaders import (
    TextLoader,
    PyMuPDFLoader,
)

import fitz

from models import ProcessingEnum
from helpers.config import get_settings
from stores.vlm.VLMProviderFactory import VLMProviderFactory
from stores.ocr.OCRProviderFactory import OCRProviderFactory


@dataclass
class Document:
    page_content: str
    metadata: dict
    image: str = None


class ProcessController(BaseController):

    def __init__(self, project_id: str):
        super().__init__()

        self.project_id = project_id

        self.project_path = (
            ProjectController().get_project_path(
                project_id=project_id
            )
        )

        settings = get_settings()

        self.logger = logging.getLogger(__name__)

        self.vlm_client = None
        self.ocr_client = None

        # ---------------------------------------------------------
        # VLM
        # ---------------------------------------------------------

        if settings.VLM_BACKEND:

            vlm_factory = VLMProviderFactory(
                config=settings
            )

            self.vlm_client = vlm_factory.create(
                provider=settings.VLM_BACKEND.lower()
            )

            if self.vlm_client:

                self.vlm_client.set_generation_model(
                    settings.VLM_MODEL_ID
                )

    # =============================================================
    # FILE HELPERS
    # =============================================================

    def get_file_extension(self, file_id: str):

        return os.path.splitext(file_id)[-1]

    # =============================================================
    # IMAGE PREPROCESSING
    # =============================================================

    def preprocess_image(
        self,
        image_bytes: bytes
    ):

        """
        Preprocess image before OCR/VLM.

        - Fix EXIF orientation
        - Convert image to RGB
        - Resize large images
        - Compress as JPEG
        - Keep aspect ratio
        - Original image is never modified
        """

        MAX_IMAGE_DIMENSION = 1536
        JPEG_QUALITY = 88
        MAX_PROCESSED_IMAGE_BYTES = 2 * 1024 * 1024

        try:

            image = Image.open(
                io.BytesIO(image_bytes)
            )

            # Fix phone-camera orientation
            image = ImageOps.exif_transpose(
                image
            )

            # Convert to RGB
            if image.mode != "RGB":

                if image.mode in (
                    "RGBA",
                    "LA",
                    "P"
                ):

                    rgba = image.convert(
                        "RGBA"
                    )

                    background = Image.new(
                        "RGB",
                        rgba.size,
                        "white"
                    )

                    background.paste(
                        rgba,
                        mask=rgba.getchannel("A")
                    )

                    image = background

                else:

                    image = image.convert(
                        "RGB"
                    )

            original_width, original_height = (
                image.size
            )

            self.logger.info(
                f"Original image size: "
                f"{original_width}x{original_height}"
            )

            # Resize only if necessary
            max_dimension = max(
                original_width,
                original_height
            )

            if max_dimension > MAX_IMAGE_DIMENSION:

                scale = (
                    MAX_IMAGE_DIMENSION
                    / max_dimension
                )

                new_width = int(
                    original_width * scale
                )

                new_height = int(
                    original_height * scale
                )

                image = image.resize(
                    (
                        new_width,
                        new_height
                    ),
                    Image.Resampling.LANCZOS
                )

                self.logger.info(
                    f"Image resized: "
                    f"{original_width}x{original_height} "
                    f"-> "
                    f"{new_width}x{new_height}"
                )

            # Compress as JPEG
            quality = JPEG_QUALITY

            while True:

                output = io.BytesIO()

                image.save(
                    output,
                    format="JPEG",
                    quality=quality,
                    optimize=True
                )

                processed_bytes = (
                    output.getvalue()
                )

                if (
                    len(processed_bytes)
                    <= MAX_PROCESSED_IMAGE_BYTES
                    or quality <= 50
                ):
                    break

                quality -= 5

            self.logger.info(
                f"Processed image: "
                f"{image.width}x{image.height}, "
                f"{len(processed_bytes) / 1024:.1f} KB, "
                f"quality={quality}"
            )

            return (
                processed_bytes,
                "image/jpeg"
            )

        except Exception as e:

            self.logger.exception(
                f"Error preprocessing image: {e}"
            )

            raise RuntimeError(
                f"Could not preprocess image: {e}"
            )

    # =============================================================
    # FILE LOADER
    # =============================================================

    def get_file_loader(
        self,
        file_id: str
    ):

        file_ext = self.get_file_extension(
            file_id=file_id
        )

        file_path = os.path.join(
            self.project_path,
            file_id
        )

        if not os.path.exists(file_path):
            return None

        if file_ext == ProcessingEnum.TXT.value:

            return TextLoader(
                file_path,
                encoding="utf-8"
            )

        if file_ext == ProcessingEnum.PDF.value:

            return PyMuPDFLoader(
                file_path
            )

        return None

    # =============================================================
    # GET FILE CONTENT
    # =============================================================

    def get_file_content(
        self,
        file_id: str,
        original_filename: str = None
    ):

        file_ext = self.get_file_extension(
            file_id=file_id
        ).lower()

        if file_ext == ProcessingEnum.PDF.value:

            return self.get_pdf_content_with_ocr(
                file_id=file_id,
                original_filename=original_filename
            )

        if file_ext in (
            ProcessingEnum.JPEG.value,
            ProcessingEnum.JPG.value,
            ProcessingEnum.PNG.value,
            ProcessingEnum.WEBP.value,
        ):

            return self.get_image_content(
                file_id=file_id,
                original_filename=original_filename
            )

        if file_ext == ProcessingEnum.XLSX.value:

            return self.get_excel_content(
                file_id=file_id,
                original_filename=original_filename
            )

        loader = self.get_file_loader(
            file_id=file_id
        )

        if loader:

            documents = loader.load()

            # -----------------------------------------------------
            # Preserve original filename for generic loaders
            # -----------------------------------------------------

            for document in documents:

                if document.metadata is None:
                    document.metadata = {}

                document.metadata[
                    "original_filename"
                ] = (
                    original_filename
                    or file_id
                )

            return documents

        return None

    # =============================================================
    # IMAGE CONTENT
    # =============================================================

    def get_image_content(
        self,
        file_id: str,
        original_filename: str = None
    ):

        try:

            file_path = Path(
                os.path.join(
                    self.project_path,
                    file_id
                )
            )

            if not file_path.exists():

                self.logger.error(
                    f"Image not found: {file_path}"
                )

                return None

            # -----------------------------------------------------
            # Read original image
            # -----------------------------------------------------

            original_bytes = (
                file_path.read_bytes()
            )

            # -----------------------------------------------------
            # Preprocess image
            # -----------------------------------------------------

            (
                processed_bytes,
                image_type
            ) = self.preprocess_image(
                image_bytes=original_bytes
            )

            image_base64 = (
                base64.b64encode(
                    processed_bytes
                ).decode("utf-8")
            )

            # -----------------------------------------------------
            # OCR
            # -----------------------------------------------------

            ocr_text = ""

            try:

                ocr_client = (
                    self.get_ocr_client()
                )

                if ocr_client:

                    ocr_result = (
                        ocr_client.extract_text(
                            image_base64=image_base64,
                            image_type=image_type
                        )
                    )

                    if ocr_result:

                        ocr_text = (
                            ocr_result.get(
                                "text",
                                ""
                            )
                            or ""
                        )

            except Exception as e:

                self.logger.exception(
                    f"OCR failed for {file_id}: {e}"
                )

            # -----------------------------------------------------
            # VLM
            # -----------------------------------------------------

            vlm_text = ""

            try:

                if self.vlm_client:

                    prompt = """
Look at this document image and identify ONLY important visual information
that OCR cannot reliably capture.

Focus on:
- document type
- tables and their structure
- stamps or seals
- signatures
- logos
- handwritten content
- diagrams or charts
- layout relationships between fields
- important visual elements

Do NOT transcribe the document text.
Do NOT repeat OCR text.
Do NOT invent or guess information.
Do NOT provide a summary.
Do NOT add an introduction or conclusion.

If there is no important visual information beyond the text,
return exactly:

NO_VISUAL_INFORMATION

Return a short plain-text answer only.
"""

                    vlm_text = (
                        self.vlm_client
                        .analyze_image_base64(
                            image_base64=image_base64,
                            prompt=prompt,
                            image_type=image_type
                        )
                        or ""
                    )

            except Exception as e:

                self.logger.error(
                    f"VLM failed for {file_id}: {e}"
                )

            # -----------------------------------------------------
            # Combine OCR + VLM
            # -----------------------------------------------------

            content_parts = []

            if ocr_text.strip():

                content_parts.append(
                    "## OCR Text\n\n"
                    + ocr_text.strip()
                )

            if vlm_text.strip():

                content_parts.append(
                    "## VLM Analysis\n\n"
                    + vlm_text.strip()
                )

            combined_content = (
                "\n\n".join(
                    content_parts
                )
            )

            if not combined_content.strip():

                self.logger.error(
                    f"No OCR/VLM content extracted "
                    f"from image: {file_id}"
                )

                return None

            # -----------------------------------------------------
            # Document
            # -----------------------------------------------------

            document = Document(
                page_content=combined_content,
                image=image_base64,
                metadata={
                    # IMPORTANT:
                    # Keep internal file_id here.
                    # It is used by filtering/RAGAS.
                    "source": file_id,

                    "original_filename": (
                        original_filename
                        or file_id
                    ),

                    "extraction_method": "ocr_vlm",

                    "has_ocr": bool(
                        ocr_text.strip()
                    ),

                    "has_vlm": bool(
                        vlm_text.strip()
                    ),

                    "image_type": image_type,

                    "original_size_bytes": len(
                        original_bytes
                    ),

                    "processed_size_bytes": len(
                        processed_bytes
                    ),
                }
            )

            return [document]

        except Exception as e:

            self.logger.exception(
                f"Error processing image "
                f"{file_id}: {e}"
            )

            return None

    # =============================================================
    # EXCEL CONTENT
    # =============================================================

    def get_excel_content(
        self,
        file_id: str,
        original_filename: str = None
    ):

        file_path = os.path.join(
            self.project_path,
            file_id
        )

        if not os.path.exists(file_path):

            self.logger.error(
                f"Excel file not found: {file_path}"
            )

            return None

        try:

            workbook = load_workbook(
                file_path,
                data_only=True
            )

            documents = []

            for sheet_name in workbook.sheetnames:

                sheet = workbook[sheet_name]

                rows = list(
                    sheet.iter_rows(
                        values_only=True
                    )
                )

                if not rows:
                    continue

                headers = rows[0]

                for row_index, row in enumerate(
                    rows[1:],
                    start=2
                ):

                    pairs = []

                    for header, value in zip(
                        headers,
                        row
                    ):

                        if value is None:
                            continue

                        header_str = (
                            str(header).strip()
                            if header is not None
                            else ""
                        )

                        pairs.append(
                            f"{header_str}: {value}"
                        )

                    if not pairs:
                        continue

                    row_text = " | ".join(
                        pairs
                    )

                    documents.append(
                        Document(
                            page_content=row_text,
                            metadata={
                                "source": file_id,

                                "original_filename": (
                                    original_filename
                                    or file_id
                                ),

                                "sheet": sheet_name,
                                "row": row_index,
                                "extraction_method": "excel",
                            }
                        )
                    )

            if not documents:

                self.logger.error(
                    f"No data extracted from "
                    f"excel file: {file_id}"
                )

                return None

            return documents

        except Exception as e:

            self.logger.exception(
                f"Error processing excel file "
                f"{file_id}: {e}"
            )

            return None

    # =============================================================
    # EXCEL STRUCTURED CONTENT
    # =============================================================

    def get_excel_structured_content(
        self,
        file_id: str
    ):

        file_path = os.path.join(
            self.project_path,
            file_id
        )

        if not os.path.exists(file_path):

            self.logger.error(
                f"Excel file not found: {file_path}"
            )

            return None

        try:

            workbook = load_workbook(
                file_path,
                data_only=True
            )

            sheets = []

            for sheet_name in workbook.sheetnames:

                sheet = workbook[sheet_name]

                rows = list(
                    sheet.iter_rows(
                        values_only=True
                    )
                )

                if not rows:
                    continue

                headers = rows[0]

                columns = []

                for header in headers:

                    if header is None:
                        continue

                    columns.append(
                        str(header).strip()
                    )

                records = []

                for row in rows[1:]:

                    record = {}

                    for index, value in enumerate(row):

                        if index >= len(headers):
                            continue

                        header = headers[index]

                        if header is None:
                            continue

                        column_name = str(
                            header
                        ).strip()

                        record[column_name] = value

                    if any(
                        value is not None
                        for value in record.values()
                    ):

                        records.append(
                            record
                        )

                sheets.append(
                    {
                        "sheet_name": sheet_name,
                        "columns": columns,
                        "row_count": len(records),
                        "records": records,
                    }
                )

            if not sheets:

                self.logger.error(
                    f"No structured data found "
                    f"in Excel file: {file_id}"
                )

                return None

            return sheets

        except Exception as e:

            self.logger.exception(
                f"Error extracting structured "
                f"Excel data {file_id}: {e}"
            )

            return None

    # =============================================================
    # IMAGE MIME TYPE
    # =============================================================

    def get_image_mime_type(
        self,
        file_id: str
    ):

        file_ext = self.get_file_extension(
            file_id=file_id
        ).lower()

        mime_types = {
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".png": "image/png",
            ".webp": "image/webp",
        }

        return mime_types.get(
            file_ext,
            "image/png"
        )

    # =============================================================
    # OCR CLIENT
    # =============================================================

    def get_ocr_client(self):

        if self.ocr_client is None:

            settings = get_settings()

            ocr_factory = OCRProviderFactory(
                config=settings
            )

            self.ocr_client = (
                ocr_factory.create(
                    provider=settings.OCR_BACKEND
                )
            )

        return self.ocr_client

    # =============================================================
    # PDF PAGE -> BASE64
    # =============================================================

    def pdf_page_to_base64(
        self,
        page
    ):

        try:

            pix = page.get_pixmap(
                matrix=fitz.Matrix(
                    2,
                    2
                ),
                alpha=False
            )

            image_bytes = pix.tobytes(
                "png"
            )

            (
                processed_bytes,
                _
            ) = self.preprocess_image(
                image_bytes=image_bytes
            )

            return base64.b64encode(
                processed_bytes
            ).decode("utf-8")

        except Exception as e:

            self.logger.exception(
                f"Error converting PDF page "
                f"to base64: {e}"
            )

            return None

    # =============================================================
    # PDF PAGE IMAGE
    # =============================================================

    def get_pdf_page_image_base64(
        self,
        file_id: str,
        page_number: int
    ):

        pdf_path = os.path.join(
            self.project_path,
            file_id
        )

        if not os.path.exists(pdf_path):
            return None

        pdf = fitz.open(
            pdf_path
        )

        try:

            page_index = page_number - 1

            if (
                page_index < 0
                or page_index >= len(pdf)
            ):
                return None

            page = pdf[page_index]

            return self.pdf_page_to_base64(
                page=page
            )

        finally:

            pdf.close()

    # =============================================================
    # PDF OCR
    # =============================================================

    def extract_pdf_page_with_ocr(
        self,
        page
    ):

        ocr_client = self.get_ocr_client()

        if not ocr_client:
            return None

        image_base64 = (
            self.pdf_page_to_base64(
                page=page
            )
        )

        result = (
            ocr_client.extract_text(
                image_base64=image_base64,
                image_type="image/jpeg"
            )
        )

        return result

    # =============================================================
    # OCR QUALITY
    # =============================================================

    def is_ocr_good(
        self,
        ocr_result: dict,
        min_avg_score: float = 0.75,
        min_chars: int = 20,
        min_words: int = 3,
    ) -> bool:

        if not ocr_result:
            return False

        text = (
            ocr_result.get("text")
            or ""
        ).strip()

        scores = (
            ocr_result.get("scores")
            or []
        )

        if not text:
            return False

        if len(text) < min_chars:
            return False

        if len(text.split()) < min_words:
            return False

        if not scores:
            return False

        try:

            scores = [
                float(score)
                for score in scores
            ]

            avg_score = (
                sum(scores)
                / len(scores)
            )

            print(
                f"[OCR] chars={len(text)}, "
                f"words={len(text.split())}, "
                f"avg_score={avg_score:.3f}"
            )

            return (
                avg_score
                >= min_avg_score
            )

        except (
            TypeError,
            ValueError
        ):

            return False

    # =============================================================
    # PDF CONTENT WITH OCR / VLM
    # =============================================================

    def get_pdf_content_with_ocr(
        self,
        file_id: str,
        original_filename: str = None
    ):

        pdf_path = os.path.join(
            self.project_path,
            file_id
        )

        if not os.path.exists(pdf_path):
            return None

        pdf = fitz.open(
            pdf_path
        )

        documents = []

        try:

            for page_number, page in enumerate(
                pdf,
                start=1
            ):

                # -------------------------------------------------
                # IMPORTANT:
                # Initialize both variables for EVERY page.
                # Native PDF pages do not execute OCR branch.
                # -------------------------------------------------

                ocr_result = None
                vlm_result = None

                # =================================================
                # 1. Native PDF text
                # =================================================

                page_text = (
                    page.get_text(
                        "text"
                    ).strip()
                )

                if page_text:

                    documents.append(
                        Document(
                            page_content=page_text,
                            metadata={
                                "page": page_number,

                                # Keep internal file ID.
                                "source": file_id,

                                "original_filename": (
                                    original_filename
                                    or file_id
                                ),

                                "extraction_method": (
                                    "pymupdf"
                                ),

                                "has_ocr": False,
                                "has_vlm": False,
                            }
                        )
                    )

                    continue

                # =================================================
                # 2. Convert page to image
                # =================================================

                image_base64 = (
                    self.pdf_page_to_base64(
                        page=page
                    )
                )

                if not image_base64:

                    self.logger.warning(
                        f"Could not convert "
                        f"PDF page {page_number} "
                        f"to image: {file_id}"
                    )

                    continue

                # =================================================
                # 3. OCR
                # =================================================

                ocr_client = (
                    self.get_ocr_client()
                )

                if ocr_client:

                    try:

                        ocr_result = (
                            ocr_client.extract_text(
                                image_base64=image_base64,
                                image_type="image/jpeg"
                            )
                        )

                    except Exception as e:

                        self.logger.exception(
                            f"OCR failed for "
                            f"{file_id}, "
                            f"page {page_number}: {e}"
                        )

                        ocr_result = None

                # =================================================
                # 4. Decide whether VLM is needed
                # =================================================

                if (
                    ocr_result
                    and self.is_ocr_good(
                        ocr_result
                    )
                ):

                    print(
                        f"[PDF][Page {page_number}] "
                        f"OCR is good -> skipping VLM"
                    )

                else:

                    print(
                        f"[PDF][Page {page_number}] "
                        f"OCR is weak -> running VLM"
                    )

                    ocr_text = ""

                    if ocr_result:

                        ocr_text = (
                            ocr_result.get(
                                "text",
                                ""
                            )
                            or ""
                        )

                    vlm_result = (
                        self.extract_pdf_page_with_vlm(
                            page=page,
                            image_base64=image_base64,
                            ocr_text=ocr_text
                        )
                    )

                # =================================================
                # 5. Combine OCR + VLM
                # =================================================

                combined_content = []

                if (
                    ocr_result
                    and ocr_result.get("text")
                ):

                    ocr_text = (
                        ocr_result["text"]
                        .strip()
                    )

                    if ocr_text:

                        combined_content.append(
                            "## OCR Content\n\n"
                            + ocr_text
                        )

                if vlm_result:

                    vlm_text = (
                        vlm_result.strip()
                    )

                    if vlm_text:

                        # Remove markdown code fences
                        if vlm_text.startswith(
                            "```markdown"
                        ):

                            vlm_text = (
                                vlm_text[
                                    len(
                                        "```markdown"
                                    ):
                                ].strip()
                            )

                        if vlm_text.endswith(
                            "```"
                        ):

                            vlm_text = (
                                vlm_text[:-3]
                                .strip()
                            )

                        combined_content.append(
                            "## VLM Analysis\n\n"
                            + vlm_text
                        )

                # =================================================
                # 6. Create multimodal document
                # =================================================

                if combined_content:

                    documents.append(
                        Document(
                            page_content=(
                                "\n\n".join(
                                    combined_content
                                )
                            ),

                            metadata={
                                "page": page_number,

                                # Keep internal ID.
                                "source": file_id,

                                "original_filename": (
                                    original_filename
                                    or file_id
                                ),

                                "extraction_method": (
                                    "ocr_vlm"
                                    if vlm_result
                                    else "ocr"
                                ),

                                "has_ocr": bool(
                                    ocr_result
                                    and ocr_result.get(
                                        "text"
                                    )
                                ),

                                "has_vlm": bool(
                                    vlm_result
                                ),
                            },

                            image=image_base64
                        )
                    )

            return documents

        finally:

            pdf.close()

    # =============================================================
    # PDF PAGE VLM
    # =============================================================

    def analyze_pdf_page_with_vlm(
        self,
        pdf_path: str,
        page_number: int = 0
    ):

        if not self.vlm_client:
            return None

        pdf = fitz.open(
            pdf_path
        )

        try:

            page = pdf[page_number]

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(
                    2,
                    2
                ),
                alpha=False
            )

            image_path = os.path.join(
                self.project_path,
                f".vlm_test_page_"
                f"{page_number + 1}.png"
            )

            pixmap.save(
                image_path
            )

            prompt = """
Analyze this document page.

Extract all readable text accurately.

Also describe important visual information such as:
- tables
- charts
- diagrams
- figures
- labels
- captions

Preserve numbers, names, headings and dates.

Do not invent information.

Return clean text suitable for a RAG system.
"""

            result = (
                self.vlm_client.analyze_image(
                    image_path=image_path,
                    prompt=prompt,
                )
            )

            if os.path.exists(
                image_path
            ):

                os.remove(
                    image_path
                )

            return result

        finally:

            pdf.close()

    # =============================================================
    # PDF PAGE VLM EXTRACTION
    # =============================================================

    def extract_pdf_page_with_vlm(
        self,
        page,
        image_base64: str = None,
        ocr_text: str = ""
    ):

        if not self.vlm_client:
            return None

        try:

            if image_base64 is None:

                image_base64 = (
                    self.pdf_page_to_base64(
                        page=page
                    )
                )

            prompt = f"""
Analyze this document page for a RAG system.

An automatic OCR system already extracted this text from the same page,
but it may contain errors:

---OCR TEXT START---
{ocr_text}
---OCR TEXT END---

Tasks:

1. Extract all readable text accurately, correcting OCR errors using the image.
2. Preserve:
- headings
- names
- dates
- numbers
- prices
- IDs
- labels
3. Reconstruct tables using Markdown tables when possible.
4. Describe important visual elements such as:
- charts
- diagrams
- figures
- forms
- signatures
- stamps
5. Preserve the logical reading order.
6. Do not invent or infer information that is not visible.

Return only the extracted and structured document content.
"""

            result = (
                self.vlm_client.analyze_image_base64(
                    image_base64=image_base64,
                    prompt=prompt,
                    image_type="image/jpeg"
                )
            )

            if result:

                return result.strip()

            return None

        except Exception as e:

            print(
                f"VLM extraction failed for page: {e}"
            )

            return None

    # =============================================================
    # SIMPLE TEXT SPLITTER
    # =============================================================

    def process_simpler_splitter(
        self,
        texts: List[str],
        metadatas: List[dict],
        chunk_size: int,
        splitter_tag: str = "\n"
    ):

        chunks = []

        for text, metadata in zip(
            texts,
            metadatas
        ):

            full_text = text or ""

            lines = [
                doc.strip()
                for doc in full_text.split(
                    splitter_tag
                )
                if len(doc.strip()) > 1
            ]

            current_chunk = ""

            for line in lines:

                current_chunk += (
                    line
                    + splitter_tag
                )

                if len(
                    current_chunk
                ) >= chunk_size:

                    chunks.append(
                        Document(
                            page_content=(
                                current_chunk.strip()
                            ),
                            metadata={
                                **metadata
                            }
                        )
                    )

                    current_chunk = ""

            if current_chunk.strip():

                chunks.append(
                    Document(
                        page_content=(
                            current_chunk.strip()
                        ),
                        metadata={
                            **metadata
                        }
                    )
                )

        return chunks

    # =============================================================
    # MULTIMODAL SPLITTER
    # =============================================================

    def process_multimodal_splitter(
        self,
        documents: List[Document],
        chunk_size: int,
        splitter_tag: str = "\n"
    ):

        sentence_end_pattern = re.compile(
            r'(?<=[.!?؟])\s+'
        )

        chunks = []

        for document in documents:

            text = (
                document.page_content.strip()
            )

            if not text:
                continue

            sentences = [
                s.strip()
                for s in sentence_end_pattern.split(
                    text
                )
                if s.strip()
            ]

            current_chunk = ""

            for sentence in sentences:

                # -------------------------------------------------
                # Sentence larger than chunk size
                # -------------------------------------------------

                if len(sentence) > chunk_size:

                    if current_chunk.strip():

                        chunks.append(
                            Document(
                                page_content=(
                                    current_chunk.strip()
                                ),
                                metadata={
                                    **document.metadata
                                },
                                image=document.image
                            )
                        )

                        current_chunk = ""

                    lines = [
                        line.strip()
                        for line in sentence.split(
                            splitter_tag
                        )
                        if len(line.strip()) > 1
                    ]

                    sub_chunk = ""

                    for line in lines:

                        if sub_chunk:
                            sub_chunk += " "

                        sub_chunk += line

                        if len(
                            sub_chunk
                        ) >= chunk_size:

                            chunks.append(
                                Document(
                                    page_content=(
                                        sub_chunk.strip()
                                    ),
                                    metadata={
                                        **document.metadata
                                    },
                                    image=document.image
                                )
                            )

                            sub_chunk = ""

                    if sub_chunk.strip():

                        chunks.append(
                            Document(
                                page_content=(
                                    sub_chunk.strip()
                                ),
                                metadata={
                                    **document.metadata
                                },
                                image=document.image
                            )
                        )

                    continue

                # -------------------------------------------------
                # Normal case
                # -------------------------------------------------

                if (
                    current_chunk
                    and
                    len(current_chunk)
                    + len(sentence)
                    + 1
                    > chunk_size
                ):

                    chunks.append(
                        Document(
                            page_content=(
                                current_chunk.strip()
                            ),
                            metadata={
                                **document.metadata
                            },
                            image=document.image
                        )
                    )

                    current_chunk = ""

                if current_chunk:

                    current_chunk += " "

                current_chunk += sentence

            if current_chunk.strip():

                chunks.append(
                    Document(
                        page_content=(
                            current_chunk.strip()
                        ),
                        metadata={
                            **document.metadata
                        },
                        image=document.image
                    )
                )

        # =========================================================
        # Merge tiny leftover chunks
        # =========================================================

        MIN_CHUNK_CHARS = 40

        merged_chunks = []

        for chunk in chunks:

            content = (
                chunk.page_content.strip()
            )

            if (
                len(content)
                < MIN_CHUNK_CHARS
                and merged_chunks
            ):

                previous = (
                    merged_chunks[-1]
                )

                previous.page_content = (
                    previous.page_content.strip()
                    + " "
                    + content
                ).strip()

                continue

            merged_chunks.append(
                chunk
            )

        return merged_chunks

    # =============================================================
    # PROCESS FILE CONTENT
    # =============================================================

    def process_file_content(
        self,
        file_content: List[Document],
        file_id: str,
        chunk_size: int,
        overlap_size: int
    ):

        if not file_content:
            return None

        file_ext = (
            self.get_file_extension(
                file_id=file_id
            ).lower()
        )

        if file_ext == ProcessingEnum.TXT.value:

            texts = [
                document.page_content
                for document in file_content
            ]

            metadatas = [
                document.metadata
                for document in file_content
            ]

            return self.process_simpler_splitter(
                texts=texts,
                metadatas=metadatas,
                chunk_size=chunk_size
            )

        if file_ext in (
            ProcessingEnum.PDF.value,
            ProcessingEnum.JPEG.value,
            ProcessingEnum.JPG.value,
            ProcessingEnum.PNG.value,
            ProcessingEnum.WEBP.value,
        ):

            return self.process_multimodal_splitter(
                documents=file_content,
                chunk_size=chunk_size
            )

        # Excel rows are already individual chunks.
        if file_ext == ProcessingEnum.XLSX.value:

            return file_content

        return None