"""
Text-to-SQL golden evaluation for the mini-rag project.

For every golden question the evaluator
  1. asks SQLController.query() (the real pipeline: LLM -> validator ->
     read-only execution),
  2. runs the question's gold_sql on the same database,
  3. compares the two RESULTS (execution accuracy). Different SQL that
     returns the same rows counts as correct.

No judge LLM is used: it is fast and free apart from the generation model.

THREE SEPARATE MODES (--mode):

  excel         Excel -> SQL tables -> questions. The "data analysis" flow:
                the Excel file is loaded into its own database and the
                golden questions are asked about it.        (default)
  ingest-check  Excel -> SQL ONLY: checks that every sheet was loaded
                correctly (row counts, sums, distinct values, nulls). No
                LLM, no questions.
  database      A database you CONNECT (no Excel, nothing is loaded or
                changed): golden questions about its tables.
                Needs --db and --questions.

Run from the src/ folder:

    PYTHONPATH=. python evaluation/sql_evaluator.py --ids S01,J03 --tag smoke
    PYTHONPATH=. python evaluation/sql_evaluator.py --tag full
    PYTHONPATH=. python evaluation/sql_evaluator.py --category arabic
    PYTHONPATH=. python evaluation/sql_evaluator.py --verify      # semantic check ON
    PYTHONPATH=. python evaluation/sql_evaluator.py --excel-name my.xlsx --questions my.json
    PYTHONPATH=. python evaluation/sql_evaluator.py --mode database --db sqlite:///assets/database/sql/ecommerce.db --questions db.json

RAGAS SQL metrics (optional, --ragas):

    PYTHONPATH=. python evaluation/sql_evaluator.py --ragas
    PYTHONPATH=. python evaluation/sql_evaluator.py --ragas --ragas-metrics datacompy,equivalence

    datacompy   : RAGAS DataCompyScore. Executes nothing new: compares the
                  predicted result table with the gold result table
                  (rows F1). No LLM. Needs: pip install datacompy
    equivalence : RAGAS LLMSQLEquivalence. A judge LLM (RAGAS_JUDGE_MODEL,
                  default qwen2.5:7b-instruct) says whether the predicted
                  SQL and the gold SQL mean the same. Slower, and a 7B judge
                  is not very reliable: use it as a second opinion only.

In excel / ingest-check mode the Excel file is NOT read from a datasets
folder. Upload it to the app first (POST /api/v1/data/upload/<project_id> or
the Upload File button), exactly like the PDF of the RAG evaluation. The
evaluator finds it by its original file name (--excel-name, or --file-id),
loads it into the project's Excel store (database_url excel://<project_id>)
and asks the questions. The database in .env, the connected databases and
the vector database are never touched.
"""

import argparse
import asyncio
import csv
import json
import os
import re
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

EVAL_DIR = Path(__file__).resolve().parent
DEFAULT_QUESTIONS = EVAL_DIR / "datasets" / "sql_golden_questions.json"
DEFAULT_EXCEL_NAME = "sql_demo.xlsx"   # original name of the uploaded file
RESULTS_DIR = EVAL_DIR / "results"

DEFAULT_PROJECT_ID = 1   # same project the RAG evaluation uses


# ==========================================================================
# Result comparison
# ==========================================================================

def _cell(value):
    """5 and 5.0 are the same answer; text is compared trimmed."""

    if value is None:
        return "NULL"

    if isinstance(value, bool):
        return str(int(value))

    if isinstance(value, (int, float)):
        return f"{round(float(value), 2):.2f}"

    return str(value).strip()


def normalize_rows(rows, columns):
    """Column names and column order are ignored; only the values count."""

    return [
        tuple(sorted(_cell(row.get(column)) for column in columns))
        for row in rows
    ]


def same_result(gold, predicted, order_matters):

    gold_rows = normalize_rows(gold["rows"], gold["columns"])
    predicted_rows = normalize_rows(predicted["rows"], predicted["columns"])

    if order_matters:
        return gold_rows == predicted_rows

    return Counter(gold_rows) == Counter(predicted_rows)


# ==========================================================================
# Core loop (no project imports: easy to test)
# ==========================================================================

async def evaluate_questions(
    questions,
    ask,
    run_gold,
    max_rows=100,
    max_retries=2,
):
    """
    ask(question, max_rows, max_retries) -> dict returned by SQLController.query
    run_gold(sql)                        -> {"columns": [...], "rows": [...]}
    """

    results = []

    for index, item in enumerate(questions, start=1):

        print(
            f"[{index}/{len(questions)}] {item['id']} | "
            f"{item['category']} | {item['question']}"
        )

        started = time.time()

        record = {
            "id": item["id"],
            "category": item["category"],
            "question": item["question"],
            "gold_sql": item.get("gold_sql"),
            "predicted_sql": None,
            "status": None,
            "attempts": None,
            "verified": None,
            "gold_rows": None,
            "predicted_rows": None,
            "error": "",
        }

        expect_error = bool(item.get("expect_error"))

        try:
            response = await ask(item["question"], max_rows, max_retries)

            record["predicted_sql"] = response.get("sql")
            record["attempts"] = response.get("attempts")
            record["verified"] = response.get("verified")

            predicted = response["result"]
            record["predicted_rows"] = predicted.get("row_count")
            record["_pred"] = predicted
            record["_order"] = bool(item.get("order_matters"))

            if expect_error:
                record["status"] = "unexpected_success"
            else:
                gold = run_gold(item["gold_sql"])
                record["gold_rows"] = len(gold["rows"])
                record["_gold"] = gold

                record["status"] = (
                    "correct"
                    if same_result(
                        gold,
                        predicted,
                        bool(item.get("order_matters")),
                    )
                    else "wrong_result"
                )

        except Exception as exc:

            record["error"] = f"{type(exc).__name__}: {str(exc)[:300]}"

            record["status"] = (
                "refused_ok" if expect_error else "sql_error"
            )

        record["seconds"] = round(time.time() - started, 1)

        print(
            f"    {record['status']} | attempts={record['attempts']} | "
            f"{record['seconds']}s | {(record['predicted_sql'] or record['error'])[:110]}"
        )

        results.append(record)

    return results


def summarize(results):

    good = {"correct", "refused_ok"}

    total = len(results)
    passed = sum(r["status"] in good for r in results)

    answerable = [r for r in results if r["status"] not in ("refused_ok", "unexpected_success")]

    by_category = defaultdict(list)
    for r in results:
        by_category[r["category"]].append(r)

    lines = [
        "",
        "=" * 80,
        "TEXT-TO-SQL RESULTS",
        "=" * 80,
        f"Execution accuracy : {passed}/{total} = {passed / total:.1%}"
        if total else "No results",
    ]

    valid = [r for r in results if r["status"] != "sql_error"]
    lines.append(f"Valid SQL rate     : {len(valid)}/{total}")

    attempts = [r["attempts"] for r in results if r["attempts"]]
    if attempts:
        lines.append(f"Avg attempts       : {sum(attempts) / len(attempts):.2f}")

    lines.append(f"Avg seconds        : {sum(r['seconds'] for r in results) / total:.1f}")
    lines.append("")
    lines.append("Per category:")

    for category, rows in by_category.items():
        ok = sum(r["status"] in good for r in rows)
        lines.append(f"  {category:14} {ok}/{len(rows)}")

    failures = [r for r in results if r["status"] not in good]

    if failures:
        lines.append("")
        lines.append("Failures:")

        for r in failures:
            lines.append(f"  {r['id']} [{r['status']}] {r['question']}")
            lines.append(f"      predicted: {r['predicted_sql'] or r['error']}")
            if r["gold_sql"]:
                lines.append(f"      gold     : {r['gold_sql']}")

    return "\n".join(lines)


# ==========================================================================
# Excel -> SQL correctness (no LLM)
# ==========================================================================

def read_excel_source(excel_path):
    """
    Reads the workbook straight with openpyxl (not with the project's loader),
    so the check does not trust the code it is checking.
    Returns {sheet_name: {"columns": [...], "rows": [{col: value}]}} using
    the same duplicate-header rule as ProcessController (Score, Score_2).
    """

    from openpyxl import load_workbook

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    sheets = {}

    for name in workbook.sheetnames:

        rows = workbook[name].iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            continue

        names, used = [], set()

        for header in headers:

            if header is None or not str(header).strip():
                names.append(None)
                continue

            base = str(header).strip()
            candidate, counter = base, 1

            while candidate.casefold() in used:
                counter += 1
                candidate = f"{base}_{counter}"

            used.add(candidate.casefold())
            names.append(candidate)

        columns = [n for n in names if n]

        if not columns:
            continue

        records = []

        for row in rows:
            record = {n: v for n, v in zip(names, row) if n}
            if any(v is not None for v in record.values()):
                records.append(record)

        sheets[name] = {"columns": columns, "rows": records}

    workbook.close()

    return sheets


def check_table(run_sql, table, columns, source_rows):
    """Returns a list of problems (empty = the table is a faithful copy)."""

    q = lambda name: '"' + name.replace('"', '""') + '"'

    problems = []

    count = run_sql(f"SELECT COUNT(*) AS n FROM {q(table)}")["rows"][0]["n"]

    if count != len(source_rows):
        problems.append(f"row count: Excel {len(source_rows)} vs SQL {count}")

    for column in columns:

        name, original, sql_type = column["name"], column["original_name"], column["type"]

        values = [r.get(original) for r in source_rows]
        non_null = [v for v in values if v is not None]

        nulls = run_sql(
            f"SELECT COUNT(*) AS n FROM {q(table)} WHERE {q(name)} IS NULL"
        )["rows"][0]["n"]

        if nulls != len(values) - len(non_null):
            problems.append(f"{original}: empty cells Excel {len(values) - len(non_null)} vs SQL {nulls}")

        if sql_type in ("INTEGER", "BIGINT", "FLOAT") and non_null:

            total = run_sql(f"SELECT SUM({q(name)}) AS s FROM {q(table)}")["rows"][0]["s"]
            expected = sum(float(v) for v in non_null)

            if total is None or abs(float(total) - expected) > 1e-6 * max(1.0, abs(expected)):
                problems.append(f"{original}: SUM Excel {expected} vs SQL {total}")

        elif sql_type == "TEXT" and non_null:

            distinct = run_sql(
                f"SELECT COUNT(DISTINCT {q(name)}) AS d FROM {q(table)}"
            )["rows"][0]["d"]

            expected = len({str(v) for v in non_null})

            if distinct != expected:
                problems.append(f"{original}: distinct values Excel {expected} vs SQL {distinct}")

    return problems


async def run_ingest_check(args, file_id):

    from controllers.ExcelSQLController import ExcelSQLController
    from controllers.ProjectController import ProjectController

    summary = await ExcelSQLController(project_id=args.project_id).ingest(
        file_id, if_exists="replace"
    )

    # the ORIGINAL uploaded file, read independently of the project's loader
    uploaded_path = os.path.join(
        ProjectController().get_project_path(project_id=args.project_id),
        file_id,
    )

    source = read_excel_source(uploaded_path)

    run_gold, provider = build_gold_runner(f"excel://{args.project_id}")

    failures = 0

    print(f"Excel sheets: {len(source)} | SQL tables created: {summary['tables_created']}")

    if len(source) != summary["tables_created"]:
        failures += 1
        print("  !! sheet count does not match table count")

    try:
        for table in summary["tables"]:

            sheet = source.get(table["sheet_name"])

            if sheet is None:
                failures += 1
                print(f"  !! {table['sheet_name']}: not found in the Excel file")
                continue

            problems = check_table(run_gold, table["table_name"], table["columns"], sheet["rows"])

            types = ", ".join(f"{c['name']}:{c['type']}" for c in table["columns"])

            print(f"{'PASS' if not problems else 'FAIL'} | {table['sheet_name']} -> {table['table_name']} "
                  f"| {len(sheet['rows'])} rows | {types}")

            for problem in problems:
                print(f"       - {problem}")

            failures += bool(problems)
    finally:
        provider.disconnect()

    print("\nExcel -> SQL check:", "ALL TABLES FAITHFUL" if not failures else f"{failures} problem(s)")

    return failures


# ==========================================================================
# RAGAS SQL metrics
# ==========================================================================

def _csv_value(value):

    if value is None:
        return ""

    if isinstance(value, bool):
        return int(value)

    if isinstance(value, float):
        return round(value, 4)

    return value


def result_to_csv(columns, rows, sort_rows):

    import io

    body = [[_csv_value(row.get(c)) for c in columns] for row in rows]

    if sort_rows:
        # DataCompy compares row by row: make the order deterministic
        body.sort(key=lambda r: [str(x) for x in r])

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(columns)
    writer.writerows(body)

    return buffer.getvalue()


def csv_pair(gold, predicted, order_matters):
    """
    Gold and predicted result tables as CSV text for RAGAS DataCompyScore.
    Column ALIASES are not errors, so predicted columns are renamed to the
    gold names when both tables have the same number of columns.
    """

    gold_columns = list(gold["columns"])
    predicted_columns = list(predicted["columns"])
    predicted_rows = predicted["rows"]

    if len(predicted_columns) == len(gold_columns):
        predicted_rows = [
            {g: row.get(p) for g, p in zip(gold_columns, predicted_columns)}
            for row in predicted_rows
        ]
        predicted_columns = gold_columns

    sort_rows = not order_matters

    return (
        result_to_csv(gold_columns, gold["rows"], sort_rows),
        result_to_csv(predicted_columns, predicted_rows, sort_rows),
    )


def strip_auto_limit(predicted_sql, gold_sql, max_rows):
    """
    The validator adds 'LIMIT <max_rows>' to every query. When the gold SQL
    has no LIMIT, that automatic LIMIT is not part of the model's answer.
    """

    import re

    if not predicted_sql or "limit" in (gold_sql or "").lower():
        return predicted_sql

    return re.sub(
        rf"\s+LIMIT\s+{max_rows}\s*;?\s*$",
        "",
        predicted_sql.strip(),
        flags=re.IGNORECASE,
    )


def schema_text(schema):

    return "\n".join(
        f"CREATE TABLE {table} ("
        + ", ".join(f"{c['name']} {c['type']}" for c in info["columns"])
        + ");"
        for table, info in schema.items()
    )


def _scores(result, metric):
    """Score column of a RAGAS result, in dataset order."""

    frame = result.to_pandas()

    column = metric.name if metric.name in frame.columns else frame.columns[-1]

    return [None if v != v else float(v) for v in frame[column].tolist()]


def run_ragas_sql_metrics(results, wanted, schema, max_rows=100):
    """
    Adds ragas_datacompy / ragas_sql_equiv to the records that returned a
    result. Failed questions are scored 0 in the summary, not sent to RAGAS.
    """

    from datasets import Dataset
    from ragas import evaluate
    from ragas.run_config import RunConfig

    for record in results:
        record["ragas_datacompy"] = None
        record["ragas_sql_equiv"] = None

    answered = [
        r for r in results
        if r["status"] in ("correct", "wrong_result") and "_gold" in r
    ]

    if not answered:
        print("RAGAS: no answered questions to score.")
        return

    run_config = RunConfig(timeout=7200, max_retries=2, max_workers=1)

    if "datacompy" in wanted:

        from ragas.metrics import DataCompyScore

        rows = []

        for r in answered:
            gold_csv, predicted_csv = csv_pair(r["_gold"], r["_pred"], r["_order"])
            rows.append({
                "user_input": r["question"],
                "response": predicted_csv,
                "reference": gold_csv,
            })

        metric = DataCompyScore()

        result = evaluate(
            dataset=Dataset.from_list(rows),
            metrics=[metric],
            run_config=run_config,
            raise_exceptions=False,
            show_progress=True,
        )

        for r, score in zip(answered, _scores(result, metric)):
            r["ragas_datacompy"] = score

    if "equivalence" in wanted:

        from ragas.llms import llm_factory
        from ragas.metrics import LLMSQLEquivalence

        import httpx
        from openai import OpenAI

        client = OpenAI(
            base_url=os.getenv("RAGAS_JUDGE_URL", "http://localhost:11434/v1"),
            api_key="ollama",
            timeout=httpx.Timeout(3600.0, connect=10.0),
            max_retries=0,
        )

        llm = llm_factory(
            model=os.getenv("RAGAS_JUDGE_MODEL", "qwen2.5:7b-instruct"),
            provider="openai",
            client=client,
            temperature=0.0,
        )

        rows = [
            {
                "user_input": r["question"],
                "response": strip_auto_limit(r["predicted_sql"], r["gold_sql"], max_rows),
                "reference": r["gold_sql"],
                "reference_contexts": [schema],
            }
            for r in answered
        ]

        metric = LLMSQLEquivalence(llm=llm)

        result = evaluate(
            dataset=Dataset.from_list(rows),
            metrics=[metric],
            llm=llm,
            run_config=run_config,
            raise_exceptions=False,
            show_progress=True,
        )

        for r, score in zip(answered, _scores(result, metric)):
            r["ragas_sql_equiv"] = score


def safe_run_ragas(results, wanted, schema, max_rows):
    """
    RAGAS is an extra: if it fails, the questions that were already answered
    (and the execution accuracy) must NOT be lost. Prints why and goes on.
    """

    try:
        run_ragas_sql_metrics(results, wanted, schema, max_rows)

    except ImportError as exc:
        print(
            f"\nRAGAS SQL metrics skipped: missing package ({exc}).\n"
            "  Try: pip install datacompy pandas  (and a ragas version that still "
            "exposes DataCompyScore / LLMSQLEquivalence in ragas.metrics)."
        )

    except Exception as exc:
        print(f"\nRAGAS SQL metrics failed: {type(exc).__name__}: {str(exc)[:300]}")
        print("  The execution-accuracy results below are still saved.")


def summarize_ragas(results):

    answerable = [r for r in results if r["category"] != "unanswerable"]

    lines = ["", "RAGAS SQL metrics:"]

    for key, label in (
        ("ragas_datacompy", "DataCompyScore (rows F1)"),
        ("ragas_sql_equiv", "LLMSQLEquivalence"),
    ):

        scored = [r[key] for r in answerable if r.get(key) is not None]

        if not scored and not any(key in r for r in results):
            continue

        if not scored:
            lines.append(f"  {label:26} no scores")
            continue

        with_failures = sum(scored) / len(answerable)

        lines.append(
            f"  {label:26} {sum(scored) / len(scored):.3f} on {len(scored)} answered"
            f" | {with_failures:.3f} over all {len(answerable)} answerable (failures = 0)"
        )

    return "\n".join(lines)


# ==========================================================================
# Project wiring
# ==========================================================================

def load_questions(path):

    data = json.loads(Path(path).read_text(encoding="utf-8"))

    for row in data:
        if not row.get("expect_error") and not row.get("gold_sql"):
            raise ValueError(f"{row.get('id')}: gold_sql is required.")

    return data


async def load_project_assets(project_id):
    """The files uploaded to the project: the same lookup the RAG evaluator uses."""

    from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
    from sqlalchemy.orm import sessionmaker

    from helpers.config import get_settings
    from models.AssetModel import AssetModel
    from models.enums.AssetTypeEnum import AssetTypeEnum

    settings = get_settings()

    engine = create_async_engine(
        f"postgresql+asyncpg://{settings.POSTGRES_USERNAME}:{settings.POSTGRES_PASSWORD}"
        f"@{settings.POSTGRES_HOST}:{settings.POSTGRES_PORT}/{settings.POSTGRES_MAIN_DATABASE}"
    )

    db_client = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    try:
        asset_model = await AssetModel.create_instance(db_client=db_client)

        assets = await asset_model.get_all_project_assets(
            asset_project_id=project_id,
            asset_type=AssetTypeEnum.FILE.value,
        )

        return [
            {
                "asset_name": a.asset_name,
                "original_filename": (a.asset_config or {}).get("original_filename") or a.asset_name,
            }
            for a in assets
        ]

    finally:
        await engine.dispose()


def _norm_name(name):
    return re.sub(r"[\s_]+", "-", (name or "").strip().lower())


def find_uploaded_excel(assets, project_id, name="", file_id=""):
    """Returns the asset (dict) of the uploaded Excel file, or explains what to do."""

    excel = [
        a for a in assets
        if a["original_filename"].lower().endswith((".xlsx", ".xlsm"))
    ]

    listing = "\n".join(
        f"    - {a['original_filename']}   (file_id: {a['asset_name']})" for a in excel
    ) or "    (none)"

    if file_id:
        found = [a for a in assets if a["asset_name"] == file_id]
    else:
        found = [a for a in excel if _norm_name(a["original_filename"]) == _norm_name(name)]

    if not found:
        raise RuntimeError(
            f"No uploaded Excel file named '{name or file_id}' in project {project_id}.\n"
            f"Upload it first (POST /api/v1/data/upload/{project_id}, or the Upload File "
            f"button), then run again. Uploaded Excel files:\n{listing}"
        )

    if len(found) > 1:
        raise RuntimeError(
            f"'{name}' was uploaded {len(found)} times in project {project_id}. "
            f"Pick one with --file-id:\n{listing}"
        )

    return found[0]


async def resolve_uploaded_excel(args):

    assets = await load_project_assets(args.project_id)

    return find_uploaded_excel(
        assets,
        args.project_id,
        name=args.excel_name,
        file_id=args.file_id,
    )


async def prepare_database(project_id, file_id):
    """Load the uploaded Excel file into the project's Excel store (SQL tables)."""

    from controllers.ExcelSQLController import ExcelSQLController

    summary = await ExcelSQLController(project_id=project_id).ingest(
        file_id,
        if_exists="replace",   # the Excel store only holds Excel data
    )

    print(
        f"Loaded {summary['tables_created']} tables / "
        f"{summary['total_rows_inserted']} rows: "
        + ", ".join(t["table_name"] for t in summary["tables"])
    )


def build_sql_controller():

    from controllers.SQLController import SQLController
    from helpers.config import get_settings
    from stores.llm.LLMProviderFactory import LLMProviderFactory
    from stores.sql.templates import SQLTemplateParser

    settings = get_settings()

    generation_client = LLMProviderFactory(settings).create(
        provider=settings.GENERATION_BACKEND
    )
    generation_client.set_generation_model(
        model_id=settings.GENERATION_MODEL_ID
    )

    return SQLController(
        generation_client=generation_client,
        sql_template_parser=SQLTemplateParser(
            language=settings.PRIMARY_LANG,
            default_language=settings.DEFAULT_LANG,
        ),
    )


def build_gold_runner(database_url=None):

    from helpers.sql_access import resolve_sql_target
    from stores.sql.SQLEnums import SQLProviderEnum
    from stores.sql.SQLProviderFactory import SQLProviderFactory

    provider_name, url = resolve_sql_target(None, database_url)

    provider = SQLProviderFactory.create(
        provider=SQLProviderEnum(provider_name),
        database_url=url,
    )

    def run_gold(sql):
        return provider.execute(sql, max_rows=10000, read_only=True)

    return run_gold, provider


async def main(args):

    if args.mode == "database":

        if not args.db:
            raise SystemExit("--mode database needs --db (the URL of the database you connect).")

        # the connected database is the one the app is configured with
        os.environ["SQL_DATABASE_URL"] = args.db
        os.environ["SQL_PROVIDER"] = "sqlite" if args.db.startswith("sqlite") else "postgresql"
        source_url = None

    else:

        # Excel data lives in its own store: nothing else is touched
        source_url = f"excel://{args.project_id}"

    os.environ["SQL_SEMANTIC_VERIFY"] = "true" if args.verify else "false"

    if args.ragas:
        unknown = {m.strip() for m in args.ragas_metrics.split(",") if m.strip()} - {"datacompy", "equivalence"}

        if unknown:
            raise SystemExit(f"Unknown --ragas-metrics: {sorted(unknown)} (use datacompy and/or equivalence)")

    uploaded = None

    if args.mode in ("excel", "ingest-check"):
        uploaded = await resolve_uploaded_excel(args)

    if args.mode == "ingest-check":
        print(f"EXCEL -> SQL CHECK | {uploaded['original_filename']} (project {args.project_id})")
        failures = await run_ingest_check(args, uploaded["asset_name"])
        raise SystemExit(1 if failures else 0)

    questions = load_questions(args.questions)

    if args.ids:
        wanted = {x.strip() for x in args.ids.split(",") if x.strip()}
        questions = [q for q in questions if q["id"] in wanted]

    if args.category:
        questions = [q for q in questions if q["category"] == args.category]

    if args.limit:
        questions = questions[:args.limit]

    if not questions:
        raise RuntimeError("No questions matched the filters.")

    print("=" * 80)
    print(f"TEXT-TO-SQL EVALUATION | tag={args.tag or 'golden'} | {len(questions)} questions")
    print(f"Mode     : {args.mode}")
    if args.mode == "database":
        print(f"Database : {args.db} (connected, read-only)")
    else:
        print(f"Excel    : {uploaded['original_filename']} (uploaded, project {args.project_id})  ->  store {source_url}")
    print(f"Verify   : {'on' if args.verify else 'off'}")
    print("=" * 80)

    if args.mode == "excel" and not args.no_ingest:
        await prepare_database(args.project_id, uploaded["asset_name"])

    controller = build_sql_controller()
    run_gold, gold_provider = build_gold_runner(source_url)

    async def ask(question, max_rows, max_retries):
        return await controller.query(
            provider=None,
            database_url=source_url,
            question=question,
            max_rows=max_rows,
            max_retries=max_retries,
        )

    try:
        results = await evaluate_questions(
            questions,
            ask=ask,
            run_gold=run_gold,
            max_rows=args.max_rows,
            max_retries=args.max_retries,
        )
    finally:
        gold_provider.disconnect()

    if args.ragas:

        wanted = {m.strip() for m in args.ragas_metrics.split(",") if m.strip()}

        schema = ""

        if "equivalence" in wanted:
            schema = schema_text(
                await controller.get_schema(provider=None, database_url=source_url)
            )

        safe_run_ragas(results, wanted, schema, args.max_rows)

    print(summarize(results))

    if args.ragas:
        print(summarize_ragas(results))

    # raw result tables are only needed for scoring
    results = [
        {k: v for k, v in r.items() if not k.startswith("_")}
        for r in results
    ]

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    base = RESULTS_DIR / f"sql_{args.tag or 'golden'}_{datetime.now():%Y%m%d_%H%M%S}"

    with open(f"{base}.json", "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    with open(f"{base}.csv", "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(dict.fromkeys(k for r in results for k in r)))
        writer.writeheader()
        writer.writerows(results)

    print(f"\nSaved: {base}.csv\nSaved: {base}.json")


def parse_args():

    parser = argparse.ArgumentParser(description="Text-to-SQL golden evaluation.")
    parser.add_argument("--mode", choices=["excel", "ingest-check", "database"], default="excel",
                        help="excel: Excel -> SQL -> questions | ingest-check: only verify the "
                             "Excel -> SQL load | database: a connected database (needs --db).")
    parser.add_argument("--questions", default=str(DEFAULT_QUESTIONS))
    parser.add_argument("--excel-name", default=DEFAULT_EXCEL_NAME,
                        help="Original file name of the Excel file you UPLOADED to the project.")
    parser.add_argument("--file-id", default="",
                        help="Internal file id (asset_name) of the uploaded Excel file, if the name is ambiguous.")
    parser.add_argument("--db", default="",
                        help="--mode database only: URL of the connected database.")
    parser.add_argument("--no-ingest", action="store_true",
                        help="Do not reload the Excel file.")
    parser.add_argument("--project-id", type=int, default=DEFAULT_PROJECT_ID)
    parser.add_argument("--ids", default="")
    parser.add_argument("--category", default="")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--tag", default="")
    parser.add_argument("--max-rows", type=int, default=100)
    parser.add_argument("--max-retries", type=int, default=2)
    parser.add_argument("--ragas", action="store_true",
                        help="Also score with RAGAS SQL metrics.")
    parser.add_argument("--ragas-metrics", default="datacompy",
                        help="datacompy and/or equivalence (comma separated).")
    parser.add_argument("--verify", action="store_true",
                        help="Turn the semantic verifier on (extra LLM call per query).")
    return parser.parse_args()


if __name__ == "__main__":
    asyncio.run(main(parse_args()))